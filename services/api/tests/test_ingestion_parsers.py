import io
import pytest
from pypdf import PdfWriter
import docx
import pptx

from app.rag.parsing.pdf_parser import PDFParser
from app.rag.parsing.docx_parser import DocxParser
from app.rag.parsing.pptx_parser import PPTXParser
from app.rag.parsing.text_parser import TextParser
from app.rag.parsing.registry import get_parser_for_filename, UnsupportedParserFormatError
from app.rag.chunking.structure_aware import StructureAwareChunker


def create_sample_pdf_bytes() -> bytes:
    """Generates a valid 2-page PDF in memory."""
    writer = PdfWriter()
    # Add page 1
    writer.add_blank_page(width=612, height=792)
    # Add page 2
    writer.add_blank_page(width=612, height=792)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def create_sample_docx_bytes() -> bytes:
    """Generates a valid DOCX document in memory with headings, paragraphs, and a table."""
    doc = docx.Document()
    doc.add_heading("Cloud Computing Architecture", level=1)
    doc.add_paragraph("Cloud computing is the on-demand delivery of IT resources over the Internet.")
    doc.add_heading("Service Models", level=2)
    doc.add_paragraph("The primary service models are IaaS, PaaS, and SaaS.")
    
    # Add a table
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Model"
    table.cell(0, 1).text = "Description"
    table.cell(1, 0).text = "IaaS"
    table.cell(1, 1).text = "Infrastructure as a Service"
    
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def create_sample_pptx_bytes() -> bytes:
    """Generates a valid 2-slide PPTX presentation in memory."""
    prs = pptx.Presentation()
    # Slide 1: Title slide
    title_slide_layout = prs.slide_layouts[0]
    slide1 = prs.slides.add_slide(title_slide_layout)
    slide1.shapes.title.text = "Introduction to Distributed Systems"
    slide1.placeholders[1].text = "Course CS 452 - Lecture 1"
    
    # Slide 2: Bullet slide
    bullet_slide_layout = prs.slide_layouts[1]
    slide2 = prs.slides.add_slide(bullet_slide_layout)
    slide2.shapes.title.text = "Key Characteristics"
    body_shape = slide2.shapes.placeholders[1]
    tf = body_shape.text_frame
    tf.text = "Concurrency of components"
    p = tf.add_paragraph()
    p.text = "Lack of a global clock"
    
    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def test_parser_registry_resolution():
    assert isinstance(get_parser_for_filename("lecture1.pdf"), PDFParser)
    assert isinstance(get_parser_for_filename("syllabus.docx"), DocxParser)
    assert isinstance(get_parser_for_filename("slides.pptx"), PPTXParser)
    assert isinstance(get_parser_for_filename("notes.txt"), TextParser)
    assert isinstance(get_parser_for_filename("README.md"), TextParser)

    with pytest.raises(UnsupportedParserFormatError):
        get_parser_for_filename("archive.zip")


def test_pdf_parsing_structure_and_pages():
    pdf_bytes = create_sample_pdf_bytes()
    parser = PDFParser()
    doc = parser.parse(pdf_bytes, "sample.pdf")

    assert doc.filename == "sample.pdf"
    assert doc.extension == ".pdf"
    assert doc.page_count == 2
    assert doc.is_scanned is True  # Blank pages have 0 chars, correctly detected as scanned/empty


def test_docx_parsing_headings_and_tables():
    docx_bytes = create_sample_docx_bytes()
    parser = DocxParser()
    doc = parser.parse(docx_bytes, "cloud.docx")

    assert doc.filename == "cloud.docx"
    assert doc.extension == ".docx"
    assert len(doc.blocks) >= 4

    headings = [b for b in doc.blocks if b.block_type == "heading"]
    assert len(headings) >= 2
    assert headings[0].content == "Cloud Computing Architecture"
    assert headings[0].section_path == "Cloud Computing Architecture"
    assert headings[1].content == "Service Models"
    assert "Cloud Computing Architecture > Service Models" in headings[1].section_path

    tables = [b for b in doc.blocks if b.block_type == "table"]
    assert len(tables) == 1
    assert "IaaS | Infrastructure as a Service" in tables[0].content


def test_pptx_parsing_slide_numbers_and_titles():
    pptx_bytes = create_sample_pptx_bytes()
    parser = PPTXParser()
    doc = parser.parse(pptx_bytes, "distributed_systems.pptx")

    assert doc.filename == "distributed_systems.pptx"
    assert doc.extension == ".pptx"
    assert doc.slide_count == 2
    assert doc.page_count == 2

    # Check slide numbers and titles
    slide1_blocks = [b for b in doc.blocks if b.slide_number == 1]
    assert any(b.slide_title == "Introduction to Distributed Systems" for b in slide1_blocks)

    slide2_blocks = [b for b in doc.blocks if b.slide_number == 2]
    assert any("Concurrency of components" in b.content for b in slide2_blocks)


def test_text_and_markdown_parsing():
    parser = TextParser()
    
    # Markdown
    md_content = """# Machine Learning 101

Supervised learning maps an input to an output based on example input-output pairs.

## Classification
- Binary classification
- Multi-class classification

```python
model.fit(X_train, y_train)
```
"""
    md_doc = parser.parse(md_content.encode("utf-8"), "ml.md")
    assert md_doc.extension == ".md"
    assert any(b.block_type == "heading" and b.content == "Machine Learning 101" for b in md_doc.blocks)
    assert any(b.block_type == "heading" and b.content == "Classification" for b in md_doc.blocks)
    assert any(b.block_type == "code" for b in md_doc.blocks)
    assert any(b.block_type == "list_item" for b in md_doc.blocks)

    # Plaintext
    txt_content = "Paragraph 1 about operating systems.\n\nParagraph 2 about memory management."
    txt_doc = parser.parse(txt_content.encode("utf-8"), "os.txt")
    assert txt_doc.extension == ".txt"
    assert len(txt_doc.blocks) == 2


def test_structure_aware_chunking_pptx_boundaries():
    pptx_bytes = create_sample_pptx_bytes()
    parser = PPTXParser()
    parsed_doc = parser.parse(pptx_bytes, "distributed_systems.pptx")

    chunker = StructureAwareChunker(target_chunk_size=500, chunk_overlap=50)
    chunks = chunker.chunk_document(parsed_doc)

    assert len(chunks) >= 2
    # Ensure chunks preserve slide numbers and never mix slide 1 and slide 2
    slide_numbers = {c.slide_number for c in chunks}
    assert 1 in slide_numbers
    assert 2 in slide_numbers

    for chunk in chunks:
        assert chunk.content_hash is not None
        assert chunk.token_count > 0
        assert chunk.slide_number in (1, 2)
        assert chunk.source_offsets.get("slide_number") is not None


def test_structure_aware_chunking_docx_section_path():
    docx_bytes = create_sample_docx_bytes()
    parser = DocxParser()
    parsed_doc = parser.parse(docx_bytes, "cloud.docx")

    chunker = StructureAwareChunker(target_chunk_size=300, min_chunk_size=20)
    chunks = chunker.chunk_document(parsed_doc)

    assert len(chunks) >= 1
    # Verify section path provenance is retained
    assert any(c.section_path is not None for c in chunks)
    for c in chunks:
        assert len(c.content.strip()) >= 20
        assert c.content_hash is not None
