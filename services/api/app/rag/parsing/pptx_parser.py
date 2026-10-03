import io
from typing import List, Optional
import pptx

from app.core.errors import AppError
from app.core.logging import logger
from app.rag.parsing.base import BaseParser, ParsedBlock, ParsedDocument


class PPTXParsingError(AppError):
    def __init__(self, message: str, code: str = "PPTX_PARSING_FAILED", status_code: int = 400):
        super().__init__(message=message, code=code, status_code=status_code)


class PPTXParser(BaseParser):
    """
    Parser for Microsoft PowerPoint (.pptx) presentations.
    Extracts slide numbers, slide titles, body text frames, and presenter notes.
    Preserves slide-level provenance for slides-to-citations mapping.
    """

    @property
    def parser_name(self) -> str:
        return "python-pptx"

    @property
    def parser_version(self) -> str:
        return getattr(pptx, "__version__", "1.0.0")

    def parse(self, file_bytes: bytes, filename: str) -> ParsedDocument:
        try:
            prs = pptx.Presentation(io.BytesIO(file_bytes))
        except Exception as exc:
            logger.error("Failed to parse PPTX presentation '%s': %s", filename, exc)
            raise PPTXParsingError(f"Corrupted or invalid PPTX presentation: {filename}")

        blocks: List[ParsedBlock] = []
        slide_count = len(prs.slides)

        for slide_idx, slide in enumerate(prs.slides):
            slide_number = slide_idx + 1  # 1-indexed

            # Extract slide title
            slide_title: Optional[str] = None
            if slide.shapes.title and slide.shapes.title.text:
                slide_title = slide.shapes.title.text.strip()

            # Iterate through shapes to extract textual content
            shape_texts: List[str] = []
            for shape in slide.shapes:
                if not shape.has_text_frame:
                    continue

                for paragraph in shape.text_frame.paragraphs:
                    line = paragraph.text.strip()
                    if line and line != slide_title:
                        shape_texts.append(line)

            # If slide has a title, emit title as a heading block
            if slide_title:
                blocks.append(
                    ParsedBlock(
                        content=slide_title,
                        block_type="heading",
                        slide_number=slide_number,
                        slide_title=slide_title,
                        heading=slide_title,
                        section_path=f"Slide {slide_number}: {slide_title}",
                    )
                )

            # Emit body text paragraphs
            for text_line in shape_texts:
                blocks.append(
                    ParsedBlock(
                        content=text_line,
                        block_type="paragraph",
                        slide_number=slide_number,
                        slide_title=slide_title,
                        heading=slide_title,
                        section_path=f"Slide {slide_number}: {slide_title}" if slide_title else f"Slide {slide_number}",
                    )
                )

            # Check for speaker notes
            try:
                if slide.has_notes_slide and slide.notes_slide:
                    notes_frame = slide.notes_slide.notes_text_frame
                    if notes_frame and notes_frame.text:
                        notes_text = notes_frame.text.strip()
                        if notes_text:
                            blocks.append(
                                ParsedBlock(
                                    content=notes_text,
                                    block_type="slide_notes",
                                    slide_number=slide_number,
                                    slide_title=slide_title,
                                    heading=slide_title,
                                    section_path=f"Slide {slide_number} (Notes)",
                                )
                            )
            except Exception as note_err:
                logger.debug("Could not read notes on slide %d: %s", slide_number, note_err)

        return ParsedDocument(
            filename=filename,
            extension=".pptx",
            page_count=slide_count,
            slide_count=slide_count,
            blocks=blocks,
            metadata={"slide_count": slide_count},
        )
